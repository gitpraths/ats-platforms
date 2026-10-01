"""
Script to update WVA CRM Testing_latest_Updated.docx and generate WVA CRM Testing_11AUG2026.docx
Adds the comprehensive non-technical updates for Testing 20/Jul/2026 & August 2026.
"""
import shutil
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC  = "/Users/deeproot/data/21MARCH2026/my-ats-platform/docs/WVA CRM Testing_latest_Updated.docx"
DEST_LATEST = "/Users/deeproot/data/21MARCH2026/my-ats-platform/docs/WVA CRM Testing_latest_Updated.docx"
DEST_AUG11  = "/Users/deeproot/data/21MARCH2026/my-ats-platform/docs/WVA CRM Testing_11AUG2026.docx"
BACKUP      = SRC + ".bak"

def tick_line(doc, label, detail):
    p = doc.add_paragraph()
    icon = p.add_run("✅ Fixed – ")
    icon.bold = True
    icon.font.color.rgb = RGBColor(0x1F, 0x7A, 0x1F)
    icon.font.size = Pt(11)
    lbl = p.add_run(label + ": ")
    lbl.bold = True
    lbl.font.size = Pt(11)
    rest = p.add_run(detail)
    rest.font.size = Pt(11)
    return p

def info_line(doc, label, detail):
    p = doc.add_paragraph()
    icon = p.add_run("ℹ️ Feature Clarification – ")
    icon.bold = True
    icon.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    icon.font.size = Pt(11)
    lbl = p.add_run(label + ": ")
    lbl.bold = True
    lbl.font.size = Pt(11)
    rest = p.add_run(detail)
    rest.font.size = Pt(11)
    return p

def bullet_line(doc, bold_label, detail):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    icon = p.add_run("  • ")
    icon.bold = True
    icon.font.size = Pt(11)
    lbl = p.add_run(bold_label + " ")
    lbl.bold = True
    lbl.font.size = Pt(11)
    rest = p.add_run(detail)
    rest.font.size = Pt(11)
    return p

def remove_existing_aug11_section(doc):
    start_idx = None
    for i, para in enumerate(doc.paragraphs):
        if "Testing 20/Jul/2026" in para.text or "Testing Updates – 20th July" in para.text:
            start_idx = i
            break
    if start_idx is None:
        return

    to_remove = [para._element for para in doc.paragraphs[start_idx:]]
    for el in to_remove:
        el.getparent().remove(el)

# Save backup
shutil.copy2(SRC, BACKUP)
print(f"Backup saved to {BACKUP}")

doc = Document(SRC)
remove_existing_aug11_section(doc)

# ── New Section Heading ───────────────────────────────────────────────────────
doc.add_paragraph() # spacer

h = doc.add_heading("📋 Testing Updates – 20th July 2026 & August 2026", level=2)
for run in h.runs:
    run.font.color.rgb = RGBColor(0x1A, 0x56, 0xAA)
    run.font.size = Pt(14)

doc.add_paragraph() # spacer

# ── Item 1: Staff Access Levels ───────────────────────────────────────────────
tick_line(
    doc,
    "Staff Access Levels & Role-Based Permissions (Admin, Staff, Training Admin)",
    "We have implemented comprehensive access control across the CRM to ensure staff members only see and manage the areas relevant to their duties, while protecting sensitive financial data and preventing unauthorized deletion of records."
)

bullet_line(
    doc,
    "1. Admin Role:",
    "Super-admin access across all 11 CRM modules, complete user management, and exclusive authorization to delete records (vacancies, candidates, placements)."
)

bullet_line(
    doc,
    "2. Staff Role:",
    "Full operational access to Dashboard, Vacancies, Candidates, Placements, Employers, Providers, and Training Program. Restricted from viewing/generating Invoices/Xero data and restricted from deleting records."
)

bullet_line(
    doc,
    "3. Training Admin Role:",
    "Dedicated access limited strictly to Candidates (view-only with full search, phone number, and email address visible) and the Training Program module. All other operational modules (Dashboard, Vacancies, Placements, Employers, Providers, Invoices) are hidden."
)

bullet_line(
    doc,
    "4. Important Updates Menu Control:",
    "The 'Important Updates' section has been restricted strictly to Admin users and hidden from non-admin navigation menus."
)

doc.add_paragraph() # spacer

# ── Item 2: Backend Audit "Who" Tracking ──────────────────────────────────────
tick_line(
    doc,
    "Backend Audit Tracking ('Who Created' & 'Who Updated')",
    "The system now automatically records the staff member's identity whenever a candidate, vacancy, placement, or application is created or modified. This audit history runs silently in the background based on the logged-in user, keeping data accurate without cluttering the screen."
)

doc.add_paragraph() # spacer

# ── Item 3: Website Vacancy Sync Clarification ────────────────────────────────
info_line(
    doc,
    "Vacancy Website Synchronization Inquiry",
    "\"Can we sync vacancies from the website? Or add in this page and it will be added to website?\""
)

bullet_line(
    doc,
    "Business Assessment & Response:",
    "Direct 2-way automatic publishing between the CRM and an external website depends on the website's Content Management System (e.g., WordPress, Webflow, or custom API)."
)

bullet_line(
    doc,
    "Current CRM Capability:",
    "Vacancies created inside WVA CRM can store job board links and export complete vacancy specification details for public listing."
)

bullet_line(
    doc,
    "Website Integration Recommendation:",
    "If automated website vacancy sync is required, an API Webhook can be connected so that clicking 'Publish Vacancy' in WVA CRM automatically posts the position live onto your website's careers page."
)

doc.add_paragraph() # spacer

# ── Item 4: Warehouse Storeperson Interview Sync ──────────────────────────────
tick_line(
    doc,
    "Warehouse Storeperson: Interview Date Pipeline Sync",
    "Setting an Interview Date now instantly updates the candidate pipeline stage to 'Interview' across both list and detail views without requiring a page refresh."
)

doc.add_paragraph() # spacer

# ── Item 5: Mobile Placement Date Clear Button ────────────────────────────────
tick_line(
    doc,
    "Mobile-Friendly Placement Date Clear Button",
    "Added a permanent 'Clear' button inside the date editor, enabling users on mobile phones, iPads, and touch devices to remove placement dates effortlessly."
)

doc.add_paragraph() # trailing spacer

# Save to both latest and AUG11 target files
doc.save(DEST_LATEST)
doc.save(DEST_AUG11)

print(f"Successfully updated:\n  {DEST_LATEST}\n  {DEST_AUG11}")
