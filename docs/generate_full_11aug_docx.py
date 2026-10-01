"""
Complete Python script to update WVA CRM Testing_latest_Updated.docx
and generate WVA CRM Testing_11AUG2026.docx.
Includes all items from Testing 20/Jul/2026 & August 2026 in plain English business language.
"""
import shutil
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC         = "/Users/deeproot/data/21MARCH2026/my-ats-platform/docs/WVA CRM Testing_latest_Updated.docx"
DEST_LATEST = "/Users/deeproot/data/21MARCH2026/my-ats-platform/docs/WVA CRM Testing_latest_Updated.docx"
DEST_AUG11  = "/Users/deeproot/data/21MARCH2026/my-ats-platform/docs/WVA CRM Testing_11AUG2026.docx"
BACKUP      = SRC + ".bak"

def section_heading(doc, text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xAA)
        run.font.size = Pt(14)
    return h

def tick_line(doc, label, detail):
    p = doc.add_paragraph()
    icon = p.add_run("✅ Fixed – ")
    icon.bold = True
    icon.font.color.rgb = RGBColor(0x1F, 0x7A, 0x1F)
    icon.font.size = Pt(11)
    lbl = p.add_run(label + ": ")
    lbl.bold = True
    lbl.font.size = Pt(11)
    rest = p.add_run(detail)
    rest.font.size = Pt(11)
    return p

def info_line(doc, label, detail):
    p = doc.add_paragraph()
    icon = p.add_run("ℹ️ Feature Clarification – ")
    icon.bold = True
    icon.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    icon.font.size = Pt(11)
    lbl = p.add_run(label + ": ")
    lbl.bold = True
    lbl.font.size = Pt(11)
    rest = p.add_run(detail)
    rest.font.size = Pt(11)
    return p

def bullet_line(doc, bold_label, detail):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    icon = p.add_run("  • ")
    icon.bold = True
    icon.font.size = Pt(11)
    lbl = p.add_run(bold_label + " ")
    lbl.bold = True
    lbl.font.size = Pt(11)
    rest = p.add_run(detail)
    rest.font.size = Pt(11)
    return p

def remove_testing_20jul_section(doc):
    start_idx = None
    for i, para in enumerate(doc.paragraphs):
        if "Testing 20/Jul/2026" in para.text or "Testing Updates – 20th July" in para.text:
            start_idx = i
            break
    if start_idx is None:
        return

    to_remove = [para._element for para in doc.paragraphs[start_idx:]]
    for el in to_remove:
        el.getparent().remove(el)

# Save backup
shutil.copy2(SRC, BACKUP)
print(f"Backup saved to {BACKUP}")

doc = Document(SRC)
remove_testing_20jul_section(doc)

# ── New Section Heading ───────────────────────────────────────────────────────
doc.add_paragraph() # spacer

section_heading(doc, "📋 Testing Updates – 20th July 2026 & August 2026")

doc.add_paragraph() # spacer

# 1. Date Validation
tick_line(
    doc,
    "Same-Day Interview & Placement Date Support",
    "The system now fully permits Interview, Training, and Placement dates to be set to the same day as the Applied/Referral date (or any date on/after applied date). This resolves scheduling issues when a candidate is interviewed or placed on the same day they were referred."
)

doc.add_paragraph()

# 2. Resume View
tick_line(
    doc,
    "Resume Viewing in New Browser Tab",
    "Clicking on a candidate's resume or document link now opens and previews the file directly in a new browser tab instead of triggering a forced background file download, providing a faster document review experience."
)

doc.add_paragraph()

# 3. Training Program Uniform Edit Form
tick_line(
    doc,
    "Uniform Training Program Edit Window",
    "The Training Program edit interface has been standardized across all training courses. Unnecessary tabs (such as redundant Certificate No tabs) have been removed, presenting a clean, consistent form for all training enrollments."
)

doc.add_paragraph()

# 4. Duplicate Candidate Redirection
tick_line(
    doc,
    "Duplicate Candidate Profile Redirection",
    "When adding a candidate, if a matching candidate (first name, last name, or email) already exists in the CRM, the system displays an alert with a direct link to open the existing candidate profile in a new tab, preventing duplicate records."
)

doc.add_paragraph()

# 5. Placement Status & Termination Record Keeping
tick_line(
    doc,
    "Placement History & Termination Reason Tracking",
    "Added dedicated fields under Placement Details to record Employment Status (Active, Terminated, Resigned, Completed), Date of Termination/Completion, and Termination Reason Notes. This preserves full candidate work history and allows adding new placement details if the candidate is re-placed in the future."
)

doc.add_paragraph()

# 6. Application History Record Keeping
tick_line(
    doc,
    "Preservation of Application History",
    "Application and pipeline history is retained even if an application is archived or removed from active pipeline boards, maintaining audit logs for compliance and reporting."
)

doc.add_paragraph()

# 7. Candidate Availability vs Benchmark Hours Separation
tick_line(
    doc,
    "Separation of Candidate Availability and Benchmark Hours",
    "Candidate Availability is now managed as a dedicated field (e.g. Monday to Friday, Weekends, Full Availability) separate from Benchmark Hours (which uses preset dropdown options: 8, 15, 20, 25, 30, 38 hours per week)."
)

doc.add_paragraph()

# 8. Training Page Cleanup
tick_line(
    doc,
    "Training Page Layout Streamlining",
    "Non-relevant benchmark fields have been removed from the Training Program page, ensuring the training interface focuses strictly on course enrollments, attendance, and completion status."
)

doc.add_paragraph()

// 9. Invoice Generation on Placement Page
tick_line(
    doc,
    "Direct Xero Invoice Access on Placement Page (Integrated with Candidate Provider)",
    "Added a direct 'Generate Invoice' action button and status indicator onto the Placement details page. When authorized staff generate a placement invoice, the system automatically resolves the candidate's linked Employment Service Provider and issues the draft invoice directly to the Provider's Xero contact, enabling seamless provider fee billing directly from the placement record."
)

doc.add_paragraph()

# 10. Placement Email Communication Options
tick_line(
    doc,
    "Placement Email Trigger Options",
    "Added direct email trigger options under Placement Details to send Employment Confirmation Emails to employers upon placement and trigger Welfare Check verification emails."
)

doc.add_paragraph()

# 11. Vacancy Sourced By / Lead Generator Tracking
tick_line(
    doc,
    "Vacancy Source Finder Tracking ('Found By Staff / Lead Generator')",
    "Added a dedicated 'Found By' field to the Vacancy creation form allowing staff to record who identified the vacancy (e.g., Staff Member, Lead Generator, or Existing Employer). This internal record is kept for internal tracking and is not displayed on public job boards."
)

doc.add_paragraph()

# 12. Staff Access Levels & Role Controls
tick_line(
    doc,
    "Staff Access Levels & Role-Based Permissions (Admin, Staff, Training Admin)",
    "We have implemented comprehensive access control across the CRM to ensure staff members only see and manage the areas relevant to their duties, while protecting sensitive financial data and preventing unauthorized deletion of records."
)

bullet_line(
    doc,
    "1. Admin Role:",
    "Super-admin access across all 11 CRM modules, complete user management, and exclusive authorization to delete records (vacancies, candidates, placements)."
)

bullet_line(
    doc,
    "2. Staff Role:",
    "Full operational access to Dashboard, Vacancies, Candidates, Placements, Employers, Providers, and Training Program. Restricted from viewing/generating Invoices/Xero data and restricted from deleting records."
)

bullet_line(
    doc,
    "3. Training Admin Role:",
    "Dedicated access limited strictly to Candidates (view-only with full search, phone number, and email address visible) and the Training Program module. All other operational modules (Dashboard, Vacancies, Placements, Employers, Providers, Invoices) are hidden."
)

bullet_line(
    doc,
    "4. Important Updates Menu Control:",
    "The 'Important Updates' section has been restricted strictly to Admin users and hidden from non-admin navigation menus."
)

doc.add_paragraph()

# 13. Backend Audit "Who" Tracking
tick_line(
    doc,
    "Backend Audit Tracking ('Who Created' & 'Who Updated')",
    "The system now automatically records the staff member's identity whenever a candidate, vacancy, placement, or application is created or modified. This audit history runs silently in the background based on the logged-in user, keeping data accurate without cluttering the screen."
)

doc.add_paragraph()

# 14. Reports Page & Excel Export Cleanups
tick_line(
    doc,
    "Report Page & Excel Export Metric Alignment",
    "Updated report card headers and Excel download sheets: renamed metrics to 'No of Referrals' for assigned jobs and interviews, removed redundant circular chart displays, and aligned report columns with active placement statistics."
)

doc.add_paragraph()

# 15. Website Vacancy Sync Clarification
info_line(
    doc,
    "Vacancy Website Synchronization Inquiry",
    "\"Can we sync vacancies from the website? Or add in this page and it will be added to website?\""
)

bullet_line(
    doc,
    "Database-Driven Website (Supported):",
    "Automated vacancy publishing is fully possible if your website is database-driven (e.g. WordPress, Webflow, dynamic CMS, or database platform). In this setup, an API Webhook can be connected so that clicking 'Publish Vacancy' in WVA CRM automatically posts the position live onto your website's careers page."
)

bullet_line(
    doc,
    "Static Website (Not Possible Directly):",
    "If your website is a static website (plain HTML without a database backend), automated live vacancy syncing is not possible directly. For static websites, staff can copy/paste the exported vacancy specifications or link to the job board URL."
)

bullet_line(
    doc,
    "Current CRM Capability:",
    "Vacancies created inside WVA CRM store job board links, manage applicant flow, and export complete vacancy specifications for public listing."
)

doc.add_paragraph() # trailing spacer

# Save to both latest and target files
doc.save(DEST_LATEST)
doc.save(DEST_AUG11)

print(f"Successfully generated updated documents:\n  {DEST_LATEST}\n  {DEST_AUG11}")
